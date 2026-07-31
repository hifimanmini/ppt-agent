#!/usr/bin/env python3
"""GPS lean system rebuild deck — SVG (1280×720) + PPTX for 向总汇报."""

from __future__ import annotations

import html
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt, Emu
from lxml import etree

ROOT = Path(__file__).resolve().parent
SLIDES_DIR = ROOT / "slides"
OUT_DIR = ROOT / "output"

# Visual direction: industrial navy + solar amber (avoid purple / cream-terracotta / broadsheet)
NAVY = "#0B2A3F"
NAVY2 = "#143D58"
AMBER = "#E07A2F"
AMBER_SOFT = "#F3E6D8"
TEAL = "#0D8A7A"
BG = "#F5F7FA"
CARD = "#FFFFFF"
TEXT = "#1A2332"
MUTED = "#5A6A7A"
LINE = "#D5DEE7"
GREEN = "#2E8B57"
RED = "#C0392B"

FONT = "DM Sans, PingFang SC, Microsoft YaHei, Noto Sans SC, sans-serif"
CJK = "PingFang SC, Microsoft YaHei, Noto Sans SC, sans-serif"


def esc(s: str) -> str:
    return html.escape(s)


def svg_wrap(body: str, bg: str = BG) -> str:
    return f'''<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">
  <defs>
    <linearGradient id="heroGrad" x1="0" y1="0" x2="1" y2="1">
      <stop offset="0%" stop-color="{NAVY}"/>
      <stop offset="100%" stop-color="{NAVY2}"/>
    </linearGradient>
    <linearGradient id="amberGrad" x1="0" y1="0" x2="1" y2="0">
      <stop offset="0%" stop-color="{AMBER}"/>
      <stop offset="100%" stop-color="#F0A05A"/>
    </linearGradient>
    <filter id="shadow" x="-5%" y="-5%" width="110%" height="120%">
      <feDropShadow dx="0" dy="2" stdDeviation="4" flood-color="#0B2A3F" flood-opacity="0.08"/>
    </filter>
  </defs>
  <rect width="1280" height="720" fill="{bg}"/>
{body}
</svg>
'''


def text(x, y, content, size=16, fill=TEXT, weight=500, anchor="start", family=None, opacity=1):
    fam = family or FONT
    return (
        f'<text x="{x}" y="{y}" fill="{fill}" font-size="{size}" font-weight="{weight}" '
        f'font-family="{fam}" text-anchor="{anchor}" opacity="{opacity}">{esc(content)}</text>'
    )


def tspan_block(x, y, lines, size=14, fill=TEXT, weight=400, lh=1.45, width=None):
    """Multi-line text using tspans."""
    parts = [
        f'<text x="{x}" y="{y}" fill="{fill}" font-size="{size}" font-weight="{weight}" '
        f'font-family="{FONT}">'
    ]
    for i, line in enumerate(lines):
        dy = 0 if i == 0 else int(size * lh)
        parts.append(f'<tspan x="{x}" dy="{dy}">{esc(line)}</tspan>')
    parts.append("</text>")
    return "\n".join(parts)


def round_rect(x, y, w, h, fill=CARD, stroke=None, r=14, shadow=False, opacity=1):
    s = f' stroke="{stroke}" stroke-width="1.5"' if stroke else ' stroke="none"'
    filt = ' filter="url(#shadow)"' if shadow else ""
    return (
        f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{r}" ry="{r}" '
        f'fill="{fill}"{s}{filt} opacity="{opacity}"/>'
    )


def badge(x, y, label, fill=AMBER):
    return (
        f'{round_rect(x, y, 48, 28, fill=fill, r=8)}'
        f'{text(x + 24, y + 19, label, size=13, fill="#fff", weight=700, anchor="middle")}'
    )


# ───────────────────── Slide builders ─────────────────────

def slide_01_cover():
    body = f'''
  <rect width="1280" height="720" fill="url(#heroGrad)"/>
  <!-- atmosphere -->
  <circle cx="1080" cy="120" r="220" fill="{AMBER}" opacity="0.08"/>
  <circle cx="200" cy="620" r="180" fill="{TEAL}" opacity="0.10"/>
  <path d="M0 520 Q320 460 640 520 T1280 500 L1280 720 L0 720 Z" fill="#081F30" opacity="0.45"/>

  {text(80, 96, "GPS 精益管理系统", size=18, fill=AMBER, weight=600)}
  {text(80, 190, "筑基 →「势起」", size=56, fill="#FFFFFF", weight=800)}
  {text(80, 250, "2026 下半年精益全域推进方案", size=28, fill="#D6E4EE", weight=500)}

  {round_rect(80, 300, 120, 4, fill=AMBER, r=2)}

  {text(80, 360, "以「453」体系框架启航精益转型", size=18, fill="#A8C5D4", weight=400)}
  {text(80, 392, "锚定「真善美」价值主张 · 打造第二生产力引擎", size=18, fill="#A8C5D4", weight=400)}

  <!-- chips -->
  {round_rect(80, 470, 140, 36, fill="rgba(255,255,255,0.08)", r=18)}
  {text(150, 493, "四级架构", size=14, fill="#fff", weight=600, anchor="middle")}
  {round_rect(236, 470, 140, 36, fill="rgba(255,255,255,0.08)", r=18)}
  {text(306, 493, "五大体系", size=14, fill="#fff", weight=600, anchor="middle")}
  {round_rect(392, 470, 140, 36, fill="rgba(255,255,255,0.08)", r=18)}
  {text(462, 493, "三阶段路径", size=14, fill="#fff", weight=600, anchor="middle")}

  {text(80, 660, "汇报对象：向总", size=14, fill="#7FA0B4", weight=500)}
  {text(1200, 660, "2026 H2", size=14, fill="#7FA0B4", weight=500, anchor="end")}
'''
    return svg_wrap(body, bg=NAVY)


def slide_02_pain():
    pains = [
        ("01", "落地不均衡", "标杆与常态两极分化",
         "标杆线成果显著，多数产线仍停在表层 5S；深层次浪费与瓶颈未触及，「重展示、轻落地、难持续」。"),
        ("02", "改善偏单点", "系统性攻坚能力不足",
         "改善多为碎片化整改，对拉晶高能耗、切片断线损耗等核心痛点攻坚不够，难以拉动良率与成本主指标。"),
        ("03", "全员氛围不足", "自主改善意识薄弱",
         "推进仍靠管理推动与专班督导，一线多为「被动执行」，尚未形成「主动找浪费、主动提改善」。"),
        ("04", "机制落地不彻底", "考核未闭环",
         "精益与人才晋升、基地评价挂钩常流于形式，缺量化指标与奖惩，管理关注度不一、基层驱动力弱。"),
        ("05", "精益信息化缺失", "改善难沉淀",
         "改善经验难沉淀、难复制，数据与现场脱节，难以支撑从运动式活动走向常态化管理。"),
    ]
    cards = []
    positions = [
        (60, 110, 370, 200),
        (455, 110, 370, 200),
        (850, 110, 370, 200),
        (60, 340, 370, 220),
        (455, 340, 370, 220),
    ]
    for (num, title, sub, desc), (x, y, w, h) in zip(pains, positions):
        cards.append(f'''
  {round_rect(x, y, w, h, fill=CARD, stroke=LINE, shadow=True)}
  {round_rect(x, y, 6, h, fill=AMBER, r=0)}
  {text(x + 24, y + 36, num, size=22, fill=AMBER, weight=800)}
  {text(x + 24, y + 68, title, size=20, fill=NAVY, weight=700)}
  {text(x + 24, y + 96, sub, size=13, fill=AMBER, weight=600)}
  {tspan_block(x + 24, y + 128, _wrap(desc, 16), size=13, fill=MUTED, lh=1.5)}
''')
    summary = _wrap("精益本质是消除浪费与持续改善；需以闭环机制与数据沉淀，破除形式主义，回归价值创造。", 22)
    body = f'''
  {text(60, 58, "当前精益推进现存真实痛点", size=32, fill=NAVY, weight=800)}
  {round_rect(60, 72, 72, 4, fill=AMBER, r=2)}
  {"".join(cards)}
  {round_rect(850, 340, 370, 220, fill=NAVY, shadow=True)}
  {text(874, 380, "关键判断", size=16, fill=AMBER, weight=700)}
  {tspan_block(874, 420, summary, size=14, fill="#E8F1F6", lh=1.55)}
'''
    return svg_wrap(body)


def _wrap(text_in: str, max_chars: int) -> list[str]:
    lines, cur = [], ""
    for ch in text_in:
        cur += ch
        if len(cur) >= max_chars and ch in "，。；、」」 ":
            lines.append(cur)
            cur = ""
        elif len(cur) >= max_chars + 4:
            lines.append(cur)
            cur = ""
    if cur:
        lines.append(cur)
    return lines or [text_in]


def slide_03_framework():
    body = f'''
  {text(60, 52, "筑基 →「势起」｜GPS 精益 453 总框架", size=28, fill=NAVY, weight=800)}
  {round_rect(60, 64, 72, 4, fill=AMBER, r=2)}

  {round_rect(60, 90, 1160, 78, fill=NAVY, r=12)}
  {tspan_block(84, 122, ["自 2025.9 起，GPS 从「探索与单点整改」迈入「体系起势与机制落地」；2026 为关键推进年，",
                          "以精益为「第二生产力」，用 453 框架践真善美，开启转型之旅。"], size=14, fill="#E8F1F6", lh=1.45)}

  <!-- 4 levels -->
  {text(60, 210, "四级架构", size=16, fill=AMBER, weight=700)}
'''
    levels = [
        ("决策层", "战略引领", "确立精益战略，审批重大课题，提供资源与政策支持"),
        ("推进层", "体系赋能", "体系规划、方法导入、项目督导、专业人才培养"),
        ("执行层", "攻坚落地", "基地一把手主责，目标分解与资源协同，确保课题出成果"),
        ("落地层", "现场改善", "班组长带一线严守 SOP，参与现场改善提案"),
    ]
    for i, (a, b, c) in enumerate(levels):
        x = 60 + i * 295
        body += f'''
  {round_rect(x, 228, 280, 130, fill=CARD, stroke=LINE, shadow=True)}
  {round_rect(x, 228, 280, 36, fill=NAVY if i == 0 else NAVY2, r=12)}
  <rect x="{x}" y="{228+24}" width="280" height="12" fill="{NAVY if i == 0 else NAVY2}"/>
  {text(x + 140, 252, f"{a} · {b}", size=14, fill="#fff", weight=700, anchor="middle")}
  {tspan_block(x + 18, 290, _wrap(c, 14), size=13, fill=MUTED, lh=1.45)}
'''
    body += f'''
  {text(60, 400, "五大体系", size=16, fill=AMBER, weight=700)}
'''
    systems = ["01 对标研学", "02 标杆复制", "03 人才育成", "04 考核管控", "05 价值生产力"]
    for i, s in enumerate(systems):
        x = 60 + i * 235
        body += f'''
  {round_rect(x, 418, 220, 64, fill=AMBER_SOFT if i % 2 == 0 else CARD, stroke=AMBER if i == 4 else LINE, shadow=True)}
  {text(x + 110, 456, s, size=15, fill=NAVY, weight=700, anchor="middle")}
'''
    body += f'''
  {text(60, 530, "三阶段路径（2026 H2）", size=16, fill=AMBER, weight=700)}
'''
    phases = [
        ("7–8 月", "筑基巩固", "补短板 · 定标准 · 全员培训"),
        ("9–10 月", "攻坚突破", "专项课题 · 核心瓶颈"),
        ("11–12 月", "价值兑现", "固化成果 · 长效机制 · 效益验证"),
    ]
    for i, (t, title, d) in enumerate(phases):
        x = 60 + i * 390
        body += f'''
  {round_rect(x, 548, 370, 110, fill=CARD, stroke=LINE, shadow=True)}
  {text(x + 24, 582, t, size=13, fill=AMBER, weight=700)}
  {text(x + 24, 612, title, size=20, fill=NAVY, weight=800)}
  {text(x + 24, 642, d, size=13, fill=MUTED, weight=500)}
'''
    return svg_wrap(body)


def slide_04_four_levels():
    rows = [
        ("01", "决策层", "战略引领", NAVY,
         ["确立公司精益战略方向与年度目标", "审批重大改善课题与跨基地资源投入", "提供核心政策支持与价值导向"]),
        ("02", "推进层", "体系赋能", NAVY2,
         ["体系规划与方法论导入", "项目督导与过程赋能", "精益专业人才育成与认证"]),
        ("03", "执行层", "攻坚落地", TEAL,
         ["基地一把手为第一责任人", "目标四级分解、资源横向协同", "确保专项课题出成果、可复制"]),
        ("04", "落地层", "现场改善", GREEN,
         ["班组长带领一线严守标准作业", "工位级浪费识别与提案改善", "5S/目视化/防错等基础动作落地"]),
    ]
    body = f'''
  {text(60, 52, "四级联动架构｜权责清晰、上下贯通", size=28, fill=NAVY, weight=800)}
  {round_rect(60, 64, 72, 4, fill=AMBER, r=2)}
  {text(60, 100, "总部决策 → 体系推进 → 基地执行 → 班组落地，杜绝工具零散、重复落地", size=15, fill=MUTED)}
'''
    for i, (num, layer, role, color, bullets) in enumerate(rows):
        y = 130 + i * 135
        body += f'''
  {round_rect(60, y, 1160, 120, fill=CARD, stroke=LINE, shadow=True)}
  {round_rect(60, y, 160, 120, fill=color, r=14)}
  <rect x="180" y="{y}" width="40" height="120" fill="{color}"/>
  {text(140, y + 52, num, size=28, fill="#fff", weight=800, anchor="middle")}
  {text(140, y + 84, layer, size=16, fill="#fff", weight=600, anchor="middle")}
  {text(230, y + 42, role, size=22, fill=NAVY, weight=800)}
'''
        for j, b in enumerate(bullets):
            body += text(230, y + 72 + j * 22, f"· {b}", size=14, fill=MUTED)
    return svg_wrap(body)


def slide_05_five_systems_overview():
    items = [
        ("01", "对标研学体系", "被动学习 → 对标赶超", "三维对标找差距，月度闭环整改"),
        ("02", "标杆复制体系", "单点标杆 → 全域标准化", "SOP 拆解复制，标杆深度升级"),
        ("03", "人才育成体系", "基础普及 → 精英赋能", "三级梯队 + 利益深度绑定"),
        ("04", "考核管控体系", "柔性要求 → 刚性闭环", "指标入考核，红黑榜闭环"),
        ("05", "价值生产力体系", "精益锚定核心课题", "拉晶/切片/组件/设备/数字化"),
    ]
    body = f'''
  {text(60, 52, "五大体系｜系统起势的主骨架", size=28, fill=NAVY, weight=800)}
  {round_rect(60, 64, 72, 4, fill=AMBER, r=2)}
'''
    for i, (num, title, sub, desc) in enumerate(items):
        y = 100 + i * 110
        accent = AMBER if i == 4 else NAVY
        body += f'''
  {round_rect(60, y, 1160, 96, fill=CARD, stroke=LINE, shadow=True)}
  {round_rect(60, y, 100, 96, fill=accent, r=14)}
  <rect x="140" y="{y}" width="20" height="96" fill="{accent}"/>
  {text(110, y + 56, num, size=28, fill="#fff", weight=800, anchor="middle")}
  {text(190, y + 40, title, size=22, fill=NAVY, weight=800)}
  {text(190, y + 72, f"{sub}  ·  {desc}", size=15, fill=MUTED)}
'''
    return svg_wrap(body)


def slide_06_systems_01_04():
    blocks = [
        ("01 对标研学", "从「被动学习」到「对标赶超」",
         ["三维对标：横向基地 + 纵向工序 + 制造标杆，识别绩效差距",
          "量化整改：能耗 / 良率 / OEE 等核心 KPI，月度差距清单与整改计划"]),
        ("02 标杆复制", "从「单点标杆」到「全域标准化」",
         ["标准化拆解：标杆车间/产线经验沉淀为 SOP 与点检标准，跨线复制",
          "标杆升级：从基础 5S 走向提质增效、降本与零缺陷，打造行业标杆"]),
        ("03 人才育成", "从「基础普及」到「精英赋能」",
         ["三级梯队：精益骨干（执行）+ 内训师（传播）+ 精益专班（决策）",
          "利益绑定：认证与晋升通道、绩效激励挂钩，形成「学会—会用—受益」"]),
        ("04 考核管控", "从「柔性要求」到「刚性闭环」",
         ["制度化考核：精益指标纳入公司正式考核，基地与核心干部量化评价",
          "红黑榜：部门负责人主责，公示改善结果，奖惩清晰、闭环管理"]),
    ]
    body = f'''
  {text(60, 48, "五大体系详解｜01–04 管理支撑系统", size=26, fill=NAVY, weight=800)}
  {round_rect(60, 60, 72, 4, fill=AMBER, r=2)}
'''
    positions = [(60, 90), (650, 90), (60, 390), (650, 390)]
    for (title, sub, bullets), (x, y) in zip(blocks, positions):
        body += f'''
  {round_rect(x, y, 570, 270, fill=CARD, stroke=LINE, shadow=True)}
  {round_rect(x, y, 570, 70, fill=NAVY, r=14)}
  <rect x="{x}" y="{y+50}" width="570" height="20" fill="{NAVY}"/>
  {text(x + 28, y + 44, title, size=20, fill="#fff", weight=800)}
  {text(x + 28, y + 105, sub, size=14, fill=AMBER, weight=600)}
'''
        for j, b in enumerate(bullets):
            wrapped = _wrap(b, 28)
            body += tspan_block(x + 28, y + 145 + j * 70, wrapped, size=14, fill=MUTED, lh=1.45)
    return svg_wrap(body)


def slide_07_value_system():
    topics = [
        ("拉晶硅损优化", "厘清硅损价值链浪费点；降低切割与磨削余量；提升硅料利用率"),
        ("切片损耗管控", "攻克断线、碎片难点；提升硅片切割良率与优品率"),
        ("组件质量攻坚", "层压缺陷、隐裂、虚焊与碎片率管控；保障可靠性与寿命"),
        ("设备效能跃升", "夯实 TPM 运行模式；提升 OEE；保障产线连续稳定运行"),
        ("精益数字化+", "以标准化为基，贯通顶层数字化模型；整合 DMS / TPMS / QMS"),
    ]
    body = f'''
  {text(60, 48, "05 价值生产力体系｜锚定核心课题、助力价值增长", size=24, fill=NAVY, weight=800)}
  {round_rect(60, 60, 72, 4, fill=AMBER, r=2)}
  {text(60, 98, "聚焦五大核心课题，突破生产瓶颈", size=16, fill=MUTED)}

  {round_rect(60, 120, 380, 430, fill="url(#heroGrad)", r=16)}
  {text(90, 180, "第二生产力", size=16, fill=AMBER, weight=600)}
  {text(90, 230, "管理 × 产能", size=32, fill="#fff", weight=800)}
  {text(90, 275, "双增长引擎", size=32, fill="#fff", weight=800)}
  {tspan_block(90, 340, _wrap("推动精益改善转化为产能增量、良率增量、效益增量，支撑业务规模化与高质量发展。", 16), size=14, fill="#A8C5D4", lh=1.5)}
'''
    for i, (title, desc) in enumerate(topics):
        y = 120 + i * 86
        body += f'''
  {round_rect(470, y, 750, 78, fill=CARD, stroke=LINE, shadow=True)}
  {round_rect(470, y, 8, 78, fill=AMBER, r=0)}
  {text(500, y + 32, f"0{i+1}  {title}", size=18, fill=NAVY, weight=800)}
  {text(500, y + 58, desc, size=13, fill=MUTED)}
'''
    body += f'''
  {round_rect(60, 570, 1160, 90, fill=AMBER_SOFT, stroke=AMBER)}
  {text(90, 610, "目标：实现管理与产能双增长，打造支撑公司高质量发展的「第二生产力」引擎", size=16, fill=NAVY, weight=700)}
  {text(90, 640, "产能增量  ·  良率增量  ·  效益增量", size=14, fill=AMBER, weight=600)}
'''
    return svg_wrap(body)


def slide_08_three_stages():
    body = f'''
  {text(60, 48, "三个阶段｜层层递进的实施路径", size=28, fill=NAVY, weight=800)}
  {round_rect(60, 60, 72, 4, fill=AMBER, r=2)}

  {round_rect(60, 85, 565, 70, fill=CARD, stroke=LINE, shadow=True)}
  {text(84, 115, "2026 H1｜筑基起势（0→1）", size=16, fill=NAVY, weight=700)}
  {text(84, 140, "建体系、树标杆、统一组织认知", size=13, fill=MUTED)}
  {round_rect(655, 85, 565, 70, fill=NAVY, shadow=True)}
  {text(679, 115, "2026 H2｜全域放大与价值兑现（1→N）", size=16, fill="#fff", weight=700)}
  {text(679, 140, "标杆复制、痛点攻坚、机制固化、产能效益兑现", size=13, fill="#A8C5D4")}

  {text(60, 190, "实施路径：三阶段推进，层层递进", size=15, fill=AMBER, weight=700)}
'''
    phases = [
        ("01", "7–8 月", "标杆认证 · 统一执行标准",
         ["标杆车间/产线认证，8 月明确复制进度，全工序夯实管理基础",
          "系统输出工具方法论，识别各基地短板，形成工序精益标准手册",
          "启动组件 DMS，实现 Andon 在线响应与问题闭环"],
         "统一执行标准，消除产线基础差距，实现标准化作业"),
        ("02", "9–10 月", "专项攻坚 · 文化激活",
         ["拉晶专项：降硅损、断线、碎片与能耗",
          "组件改善周加速；黑带人才聚焦日常管理、TPM、品质",
          "人才赋能落地，绩效与改善成果、人才培养硬挂钩"],
         "攻克生产瓶颈，激发组织内生动力，形成全员参与氛围"),
        ("03", "11–12 月", "成果固化 · 价值兑现",
         ["半年度精益人才认证与项目成果发布",
          "全年改善复盘，建设标准体系与最佳实践案例库",
          "量化降本提质增效经济价值，完成价值闭环"],
         "形成长效精益管理机制，管理成果固化、价值清晰可见"),
    ]
    for i, (num, time, theme, bullets, goal) in enumerate(phases):
        x = 60 + i * 395
        body += f'''
  {round_rect(x, 215, 375, 440, fill=CARD, stroke=LINE, shadow=True)}
  {round_rect(x, 215, 375, 88, fill=NAVY if i != 1 else AMBER, r=14)}
  <rect x="{x}" y="{215+68}" width="375" height="20" fill="{NAVY if i != 1 else AMBER}"/>
  {text(x + 24, 250, f"{num}  {time}", size=14, fill=AMBER if i != 1 else "#fff", weight=700)}
  {text(x + 24, 280, theme, size=17, fill="#fff", weight=800)}
'''
        yy = 330
        for b in bullets:
            wrapped = _wrap("· " + b, 18)
            body += tspan_block(x + 20, yy, wrapped, size=13, fill=MUTED, lh=1.4)
            yy += 18 + len(wrapped) * 18
        body += f'''
  {round_rect(x + 16, 560, 343, 72, fill=AMBER_SOFT, r=10)}
  {tspan_block(x + 28, 588, _wrap("目标：" + goal, 18), size=12, fill=NAVY, weight=600, lh=1.4)}
'''
    return svg_wrap(body)


def slide_09_truth_good_beauty():
    cols = [
        ("守「真」", "可信改善", "数据驱动、过程可控",
         ["全流程可追溯：损耗/OEE 实时采集，量化收益",
          "SPC 底线：工艺参数设限，异常自动预警",
          "根因分析：拒绝经验拍脑袋，以真实缺陷为依据"],
         "客观数据支撑，拒绝假成果，让改善可见"),
        ("行「善」", "有效改善", "守正创新、系统攻坚",
         ["夯实基础：TPM / SMED，闭环到一线",
          "攻坚痛点：拉晶、切片、组件核心瓶颈全员改善",
          "一地一策：适配各基地特征定制方案"],
         "系统突破核心痛点，让改善有价值"),
        ("求「美」", "长效改善", "极致现场、零缺陷品质",
         ["现场之美：全域 6S 与目视化，消除隐形浪费",
          "品质之美：防错、巡检、EL，严防不良流出",
          "零缺陷：批次稳定，降低客诉损失"],
         "成果标准化、持续精进，让改善成日常"),
    ]
    body = f'''
  {text(60, 48, "精益全域落地｜以「真善美」锚定改善底层逻辑", size=24, fill=NAVY, weight=800)}
  {round_rect(60, 60, 72, 4, fill=AMBER, r=2)}
  {text(60, 95, "所有精益课题、现场改善、体系运行统一对标公司核心价值主张，构建可信、有效、长效的改善闭环", size=13, fill=MUTED)}
'''
    for i, (title, outcome, defn, bullets, foot) in enumerate(cols):
        x = 60 + i * 395
        body += f'''
  {round_rect(x, 120, 375, 520, fill=CARD, stroke=LINE, shadow=True)}
  {round_rect(x, 120, 375, 110, fill=NAVY if i != 1 else TEAL, r=14)}
  <rect x="{x}" y="{120+90}" width="375" height="20" fill="{NAVY if i != 1 else TEAL}"/>
  {text(x + 28, 165, title, size=28, fill="#fff", weight=800)}
  {text(x + 28, 200, defn, size=14, fill="#A8C5D4")}
'''
        for j, b in enumerate(bullets):
            body += tspan_block(x + 24, 270 + j * 70, _wrap("· " + b, 18), size=14, fill=MUTED, lh=1.4)
        body += f'''
  {round_rect(x + 16, 520, 343, 96, fill=AMBER_SOFT if i != 1 else "#E6F5F2", r=12)}
  {text(x + 32, 555, outcome, size=16, fill=AMBER if i != 1 else TEAL, weight=800)}
  {tspan_block(x + 32, 585, _wrap(foot, 18), size=13, fill=NAVY, lh=1.4)}
'''
    return svg_wrap(body)


def slide_10_close():
    body = f'''
  <rect width="1280" height="720" fill="url(#heroGrad)"/>
  <circle cx="1100" cy="160" r="200" fill="{AMBER}" opacity="0.10"/>
  <circle cx="180" cy="580" r="160" fill="{TEAL}" opacity="0.12"/>

  {text(80, 120, "下一步", size=18, fill=AMBER, weight=600)}
  {text(80, 190, "从「势起」到「势成」", size=44, fill="#fff", weight=800)}
  {text(80, 250, "以四级联动贯通组织，以五大体系固化机制，以三阶段兑现价值", size=18, fill="#A8C5D4")}

  {round_rect(80, 310, 340, 160, fill="rgba(255,255,255,0.08)", r=16)}
  {text(110, 360, "统一标准", size=20, fill="#fff", weight=700)}
  {text(110, 400, "工具、模板、考核一把尺子", size=14, fill="#A8C5D4")}

  {round_rect(450, 310, 340, 160, fill="rgba(255,255,255,0.08)", r=16)}
  {text(480, 360, "攻坚课题", size=20, fill="#fff", weight=700)}
  {text(480, 400, "硅损 · 切片 · 组件 · OEE · 数字化", size=14, fill="#A8C5D4")}

  {round_rect(820, 310, 340, 160, fill="rgba(255,255,255,0.08)", r=16)}
  {text(850, 360, "长效闭环", size=20, fill="#fff", weight=700)}
  {text(850, 400, "真数据 · 善攻坚 · 美现场", size=14, fill="#A8C5D4")}

  {text(80, 560, "请向总审阅指导", size=22, fill="#fff", weight=700)}
  {text(80, 600, "GPS 精益管理系统  ·  2026 H2", size=14, fill="#7FA0B4")}
'''
    return svg_wrap(body, bg=NAVY)


SLIDE_BUILDERS = [
    ("slide-01.svg", "封面 — 筑基→势起", slide_01_cover),
    ("slide-02.svg", "痛点诊断", slide_02_pain),
    ("slide-03.svg", "453 总框架", slide_03_framework),
    ("slide-04.svg", "四级联动架构", slide_04_four_levels),
    ("slide-05.svg", "五大体系总览", slide_05_five_systems_overview),
    ("slide-06.svg", "体系详解 01–04", slide_06_systems_01_04),
    ("slide-07.svg", "价值生产力体系", slide_07_value_system),
    ("slide-08.svg", "三阶段路径", slide_08_three_stages),
    ("slide-09.svg", "真善美底层逻辑", slide_09_truth_good_beauty),
    ("slide-10.svg", "下一步", slide_10_close),
]


def write_svgs():
    SLIDES_DIR.mkdir(parents=True, exist_ok=True)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    labels = []
    for name, label, fn in SLIDE_BUILDERS:
        content = fn()
        (SLIDES_DIR / name).write_text(content, encoding="utf-8")
        (OUT_DIR / name).write_text(content, encoding="utf-8")
        labels.append({"file": name, "label": label})
        print("wrote", name)
    return labels


def write_html(labels):
    template = (ROOT.parents[1] / "skills/_shared/assets/preview-template.html").read_text(encoding="utf-8")
    import json
    html_out = (
        template.replace("{{TITLE}}", "GPS精益｜筑基→势起")
        .replace("{{LOGO}}", "GPS")
        .replace("{{ACCENT_COLOR}}", AMBER)
        .replace("{{SLIDES_JSON}}", json.dumps(labels, ensure_ascii=False))
    )
    (OUT_DIR / "index.html").write_text(html_out, encoding="utf-8")
    print("wrote index.html")


def set_run_font(run, size=12, bold=False, color=RGBColor(0x1A, 0x23, 0x32), name="Microsoft YaHei"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    ea.set("typeface", name)


def add_box(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    return sh


def add_tb(slide, l, t, w, h, text, size=14, bold=False, color=RGBColor(0x1A, 0x23, 0x32), align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=bold, color=color)
    return box


def build_pptx():
    """Compact PPTX mirror of the 10-slide narrative for editable delivery."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    C_NAVY = RGBColor(0x0B, 0x2A, 0x3F)
    C_AMBER = RGBColor(0xE0, 0x7A, 0x2F)
    C_BG = RGBColor(0xF5, 0xF7, 0xFA)
    C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    C_MUTED = RGBColor(0x5A, 0x6A, 0x7A)
    C_TEXT = RGBColor(0x1A, 0x23, 0x32)
    C_SOFT = RGBColor(0xF3, 0xE6, 0xD8)
    C_LINE = RGBColor(0xD5, 0xDE, 0xE7)
    C_TEAL = RGBColor(0x0D, 0x8A, 0x7A)

    # 1 Cover
    s = prs.slides.add_slide(blank)
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, C_NAVY)
    add_tb(s, Inches(0.7), Inches(1.0), Inches(10), Inches(0.4), "GPS 精益管理系统", 16, True, C_AMBER)
    add_tb(s, Inches(0.7), Inches(1.6), Inches(11), Inches(0.8), "筑基 →「势起」", 40, True, C_WHITE)
    add_tb(s, Inches(0.7), Inches(2.5), Inches(11), Inches(0.5), "2026 下半年精益全域推进方案", 22, False, RGBColor(0xD6, 0xE4, 0xEE))
    add_tb(s, Inches(0.7), Inches(3.4), Inches(11), Inches(0.8),
           "以「453」体系框架启航精益转型\n锚定「真善美」价值主张 · 打造第二生产力引擎", 16, False, RGBColor(0xA8, 0xC5, 0xD4))
    add_tb(s, Inches(0.7), Inches(6.6), Inches(6), Inches(0.3), "汇报对象：向总", 12, False, RGBColor(0x7F, 0xA0, 0xB4))

    # 2 Pain
    s = prs.slides.add_slide(blank)
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, C_BG)
    add_tb(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.5), "当前精益推进现存真实痛点", 26, True, C_NAVY)
    pains = [
        ("01 落地不均衡", "标杆与常态两极分化：多数产线停在表层 5S，深层次浪费未触及。"),
        ("02 改善偏单点", "对拉晶能耗、切片断线等核心痛点攻坚不足，难拉动良率与成本。"),
        ("03 全员氛围不足", "推进靠管理推动，一线被动执行，自主改善文化未形成。"),
        ("04 机制未闭环", "与晋升/基地评价挂钩流于形式，缺量化指标与奖惩。"),
        ("05 信息化缺失", "改善难沉淀、难复制，难以从运动式走向常态化。"),
    ]
    for i, (t, d) in enumerate(pains):
        col, row = i % 3, i // 3
        if i == 4:
            col, row = 1, 1
        x = Inches(0.45 + col * 4.2)
        y = Inches(1.1 + row * 2.7)
        add_box(s, x, y, Inches(4.0), Inches(2.4), C_WHITE, C_LINE)
        add_tb(s, x + Inches(0.2), y + Inches(0.25), Inches(3.6), Inches(0.4), t, 16, True, C_AMBER)
        add_tb(s, x + Inches(0.2), y + Inches(0.8), Inches(3.6), Inches(1.4), d, 13, False, C_MUTED)
    add_box(s, Inches(8.85), Inches(3.8), Inches(4.0), Inches(2.4), C_NAVY)
    add_tb(s, Inches(9.1), Inches(4.1), Inches(3.5), Inches(0.4), "关键判断", 14, True, C_AMBER)
    add_tb(s, Inches(9.1), Inches(4.6), Inches(3.5), Inches(1.3),
           "需以闭环机制与数据沉淀，破除形式主义，从表层整改走向深层次攻坚与价值创造。", 13, False, C_WHITE)

    # 3 Framework
    s = prs.slides.add_slide(blank)
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, C_BG)
    add_tb(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.45), "筑基→势起｜GPS 精益 453 总框架", 24, True, C_NAVY)
    add_box(s, Inches(0.5), Inches(0.95), Inches(12.3), Inches(0.9), C_NAVY)
    add_tb(s, Inches(0.7), Inches(1.1), Inches(12), Inches(0.6),
           "2026 为关键推进年：以精益为「第二生产力」，用四级架构×五大体系×三阶段路径践真善美。", 14, False, C_WHITE)
    for i, (a, b) in enumerate([("决策层", "战略引领"), ("推进层", "体系赋能"), ("执行层", "攻坚落地"), ("落地层", "现场改善")]):
        x = Inches(0.5 + i * 3.2)
        add_box(s, x, Inches(2.2), Inches(3.0), Inches(1.4), C_WHITE, C_LINE)
        add_tb(s, x + Inches(0.15), Inches(2.45), Inches(2.7), Inches(0.4), a, 16, True, C_NAVY)
        add_tb(s, x + Inches(0.15), Inches(2.95), Inches(2.7), Inches(0.4), b, 14, False, C_AMBER)
    for i, name in enumerate(["01 对标研学", "02 标杆复制", "03 人才育成", "04 考核管控", "05 价值生产力"]):
        x = Inches(0.5 + i * 2.5)
        add_box(s, x, Inches(4.0), Inches(2.35), Inches(0.8), C_SOFT if i == 4 else C_WHITE, C_AMBER if i == 4 else C_LINE)
        add_tb(s, x, Inches(4.2), Inches(2.35), Inches(0.4), name, 13, True, C_NAVY, PP_ALIGN.CENTER)
    for i, (t, d) in enumerate([("7–8月 筑基巩固", "补短板·定标准·全员培训"), ("9–10月 攻坚突破", "专项课题·核心瓶颈"), ("11–12月 价值兑现", "固化成果·效益验证")]):
        x = Inches(0.5 + i * 4.2)
        add_box(s, x, Inches(5.2), Inches(4.0), Inches(1.6), C_WHITE, C_LINE)
        add_tb(s, x + Inches(0.2), Inches(5.45), Inches(3.6), Inches(0.4), t, 15, True, C_NAVY)
        add_tb(s, x + Inches(0.2), Inches(5.95), Inches(3.6), Inches(0.5), d, 13, False, C_MUTED)

    # 4 Four levels
    s = prs.slides.add_slide(blank)
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, C_BG)
    add_tb(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.45), "四级联动架构｜权责清晰、上下贯通", 24, True, C_NAVY)
    levels = [
        (C_NAVY, "01 决策层 · 战略引领", "确立战略、审批重大课题、提供资源与政策支持"),
        (RGBColor(0x14, 0x3D, 0x58), "02 推进层 · 体系赋能", "体系规划、方法导入、项目督导、人才培养"),
        (C_TEAL, "03 执行层 · 攻坚落地", "基地一把手主责，目标分解与资源协同，确保出成果"),
        (RGBColor(0x2E, 0x8B, 0x57), "04 落地层 · 现场改善", "班组严守 SOP，工位浪费识别与全员改善提案"),
    ]
    for i, (c, t, d) in enumerate(levels):
        y = Inches(1.1 + i * 1.45)
        add_box(s, Inches(0.5), y, Inches(12.3), Inches(1.3), C_WHITE, C_LINE)
        add_box(s, Inches(0.5), y, Inches(0.25), Inches(1.3), c)
        add_tb(s, Inches(1.0), y + Inches(0.25), Inches(11), Inches(0.4), t, 18, True, C_NAVY)
        add_tb(s, Inches(1.0), y + Inches(0.7), Inches(11), Inches(0.4), d, 14, False, C_MUTED)

    # 5 Systems overview
    s = prs.slides.add_slide(blank)
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, C_BG)
    add_tb(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.45), "五大体系｜系统起势的主骨架", 24, True, C_NAVY)
    systems = [
        "01 对标研学：被动学习 → 对标赶超；三维对标 + 月度整改闭环",
        "02 标杆复制：单点标杆 → 全域标准化；SOP 拆解复制与标杆升级",
        "03 人才育成：基础普及 → 精英赋能；三级梯队 + 利益深度绑定",
        "04 考核管控：柔性要求 → 刚性闭环；指标入考核 + 红黑榜",
        "05 价值生产力：以精益锚定核心课题，突破生产瓶颈、助力价值增长",
    ]
    for i, line in enumerate(systems):
        y = Inches(1.0 + i * 1.15)
        add_box(s, Inches(0.5), y, Inches(12.3), Inches(1.0), C_WHITE, C_LINE)
        add_box(s, Inches(0.5), y, Inches(0.25), Inches(1.0), C_AMBER if i == 4 else C_NAVY)
        add_tb(s, Inches(1.0), y + Inches(0.3), Inches(11.5), Inches(0.5), line, 16, True, C_NAVY)

    # 6 Systems 01-04 detail
    s = prs.slides.add_slide(blank)
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, C_BG)
    add_tb(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.4), "五大体系详解｜01–04 管理支撑系统", 22, True, C_NAVY)
    details = [
        ("01 对标研学", "三维对标找差距；能耗/良率/OEE 月度差距清单与整改计划。"),
        ("02 标杆复制", "标杆经验沉淀 SOP 跨线复制；从 5S 升级到提质降本零缺陷。"),
        ("03 人才育成", "骨干 + 内训师 + 专班三级梯队；认证与晋升/激励挂钩。"),
        ("04 考核管控", "精益指标入正式考核；负责人主责，红黑榜公示闭环。"),
    ]
    for i, (t, d) in enumerate(details):
        col, row = i % 2, i // 2
        x, y = Inches(0.5 + col * 6.4), Inches(1.0 + row * 3.0)
        add_box(s, x, y, Inches(6.1), Inches(2.7), C_WHITE, C_LINE)
        add_box(s, x, y, Inches(6.1), Inches(0.7), C_NAVY)
        add_tb(s, x + Inches(0.25), y + Inches(0.2), Inches(5.5), Inches(0.4), t, 16, True, C_WHITE)
        add_tb(s, x + Inches(0.25), y + Inches(1.1), Inches(5.5), Inches(1.3), d, 14, False, C_MUTED)

    # 7 Value system
    s = prs.slides.add_slide(blank)
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, C_BG)
    add_tb(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.4), "05 价值生产力体系｜五大核心课题", 22, True, C_NAVY)
    add_box(s, Inches(0.5), Inches(0.9), Inches(4.0), Inches(5.5), C_NAVY)
    add_tb(s, Inches(0.8), Inches(1.4), Inches(3.4), Inches(0.4), "第二生产力", 14, True, C_AMBER)
    add_tb(s, Inches(0.8), Inches(2.0), Inches(3.4), Inches(1.2), "管理 × 产能\n双增长引擎", 26, True, C_WHITE)
    add_tb(s, Inches(0.8), Inches(3.6), Inches(3.4), Inches(1.5),
           "转化为产能增量、良率增量、效益增量，支撑规模化高质量发展。", 13, False, RGBColor(0xA8, 0xC5, 0xD4))
    topics = [
        "01 拉晶硅损优化 — 厘清浪费点，提升硅料利用率",
        "02 切片损耗管控 — 攻克断线碎片，提升优品率",
        "03 组件质量攻坚 — 隐裂/虚焊/层压缺陷与可靠性",
        "04 设备效能跃升 — TPM + OEE，保障连续稳定",
        "05 精益数字化+ — 贯通 DMS / TPMS / QMS",
    ]
    for i, line in enumerate(topics):
        y = Inches(0.9 + i * 1.1)
        add_box(s, Inches(4.8), y, Inches(8.0), Inches(0.95), C_WHITE, C_LINE)
        add_tb(s, Inches(5.1), y + Inches(0.28), Inches(7.5), Inches(0.5), line, 15, True, C_NAVY)

    # 8 Three stages
    s = prs.slides.add_slide(blank)
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, C_BG)
    add_tb(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.4), "三个阶段｜层层递进的实施路径", 22, True, C_NAVY)
    stages = [
        ("01  7–8月", "标杆认证·统一标准", "认证复制、方法论手册、组件 DMS/Andon 闭环", "统一执行标准，消除基础差距"),
        ("02  9–10月", "专项攻坚·文化激活", "拉晶专项、组件改善周、黑带赋能与绩效挂钩", "攻克瓶颈，激发全员参与"),
        ("03  11–12月", "成果固化·价值兑现", "人才认证、案例库、量化经济价值闭环", "长效机制，价值清晰可见"),
    ]
    for i, (t, theme, acts, goal) in enumerate(stages):
        x = Inches(0.45 + i * 4.25)
        add_box(s, x, Inches(1.0), Inches(4.05), Inches(5.8), C_WHITE, C_LINE)
        add_box(s, x, Inches(1.0), Inches(4.05), Inches(1.3), C_AMBER if i == 1 else C_NAVY)
        add_tb(s, x + Inches(0.2), Inches(1.2), Inches(3.6), Inches(0.35), t, 14, True, C_WHITE)
        add_tb(s, x + Inches(0.2), Inches(1.65), Inches(3.6), Inches(0.4), theme, 16, True, C_WHITE)
        add_tb(s, x + Inches(0.2), Inches(2.6), Inches(3.6), Inches(2.2), acts, 14, False, C_MUTED)
        add_box(s, x + Inches(0.2), Inches(5.5), Inches(3.65), Inches(1.0), C_SOFT)
        add_tb(s, x + Inches(0.35), Inches(5.7), Inches(3.4), Inches(0.7), "目标：" + goal, 13, True, C_NAVY)

    # 9 真善美
    s = prs.slides.add_slide(blank)
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, C_BG)
    add_tb(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.4), "精益全域落地｜以「真善美」锚定改善底层逻辑", 20, True, C_NAVY)
    cols = [
        (C_NAVY, "守「真」· 可信改善", "全流程追溯、SPC 底线、根因分析\n拒绝假改善，让改善可见"),
        (C_TEAL, "行「善」· 有效改善", "TPM/SMED 夯基、痛点攻坚、一地一策\n系统突破，让改善有价值"),
        (C_NAVY, "求「美」· 长效改善", "6S 目视化、防错/EL、零缺陷\n成果固化，让改善成日常"),
    ]
    for i, (c, t, d) in enumerate(cols):
        x = Inches(0.5 + i * 4.25)
        add_box(s, x, Inches(1.1), Inches(4.05), Inches(5.5), C_WHITE, C_LINE)
        add_box(s, x, Inches(1.1), Inches(4.05), Inches(1.2), c)
        add_tb(s, x + Inches(0.25), Inches(1.45), Inches(3.5), Inches(0.5), t, 18, True, C_WHITE)
        add_tb(s, x + Inches(0.25), Inches(2.7), Inches(3.5), Inches(3.2), d, 15, False, C_MUTED)

    # 10 Close
    s = prs.slides.add_slide(blank)
    add_box(s, 0, 0, prs.slide_width, prs.slide_height, C_NAVY)
    add_tb(s, Inches(0.7), Inches(1.3), Inches(11), Inches(0.4), "下一步", 16, True, C_AMBER)
    add_tb(s, Inches(0.7), Inches(1.9), Inches(11), Inches(0.7), "从「势起」到「势成」", 36, True, C_WHITE)
    add_tb(s, Inches(0.7), Inches(2.8), Inches(11), Inches(0.5),
           "以四级联动贯通组织，以五大体系固化机制，以三阶段兑现价值", 16, False, RGBColor(0xA8, 0xC5, 0xD4))
    for i, (t, d) in enumerate([("统一标准", "工具、模板、考核一把尺子"), ("攻坚课题", "硅损·切片·组件·OEE·数字化"), ("长效闭环", "真数据·善攻坚·美现场")]):
        x = Inches(0.7 + i * 4.1)
        add_box(s, x, Inches(3.8), Inches(3.8), Inches(1.6), RGBColor(0x14, 0x3D, 0x58))
        add_tb(s, x + Inches(0.3), Inches(4.1), Inches(3.2), Inches(0.4), t, 18, True, C_WHITE)
        add_tb(s, x + Inches(0.3), Inches(4.65), Inches(3.2), Inches(0.4), d, 13, False, RGBColor(0xA8, 0xC5, 0xD4))
    add_tb(s, Inches(0.7), Inches(6.2), Inches(10), Inches(0.4), "请向总审阅指导", 18, True, C_WHITE)

    out = OUT_DIR / "GPS精益_筑基势起_向总汇报.pptx"
    prs.save(str(out))
    print("wrote", out)
    return out


def write_speaker_notes():
    notes = """# Speaker Notes: GPS精益｜筑基→势起

## Slide 01: 封面
- 开场点明：2026 下半年进入「势起」阶段，汇报对象向总
- Time: ~30s

## Slide 02: 痛点
- 先共情五类真实痛点，再引出「必须体系化」的必要性
- Time: ~2min

## Slide 03–05: 453 框架与四级/五大
- 强调权责贯通与体系骨架，避免工具散落
- Time: ~4min

## Slide 06–07: 体系详解与五大课题
- 课题对齐拉晶/切片/组件/设备/数字化，便于拍板资源
- Time: ~4min

## Slide 08: 三阶段
- 给出 7–12 月可检查的里程碑
- Time: ~3min

## Slide 09: 真善美
- 价值主张落地为可信/有效/长效三条闭环
- Time: ~2min

## Slide 10: 下一步
- 请向总审阅：标准、课题、闭环三项请示
- Time: ~1min

**Total estimated time:** ~16–18 minutes
"""
    (OUT_DIR / "speaker-notes.md").write_text(notes, encoding="utf-8")


def write_readme():
    (ROOT / "README.md").write_text(
        """# GPS 精益｜筑基→「势起」（向总汇报重建版）

基于《向总 - 副本》原稿内容，按 ppt-agent 规范重做的 10 页演示稿。

## 交付物

| 文件 | 说明 |
|------|------|
| `output/slide-01.svg` … `slide-10.svg` | 1280×720 设计稿 |
| `output/index.html` | 交互预览（Gallery / Scroll / Present） |
| `output/GPS精益_筑基势起_向总汇报.pptx` | 可编辑 PPTX |
| `output/speaker-notes.md` | 演讲备注 |
| `generate_deck.py` | 一键重生脚本 |

## 页面结构

1. 封面：筑基→势起
2. 当前精益推进现存真实痛点
3. 453 总框架
4. 四级联动架构
5. 五大体系总览
6. 体系详解 01–04
7. 价值生产力体系（五大核心课题）
8. 三阶段实施路径
9. 真善美底层逻辑
10. 下一步

## 预览

```bash
open output/index.html   # 或 xdg-open
python generate_deck.py  # 重新生成
```
""",
        encoding="utf-8",
    )


def main():
    labels = write_svgs()
    write_html(labels)
    write_speaker_notes()
    write_readme()
    build_pptx()
    print("DONE")


if __name__ == "__main__":
    main()
