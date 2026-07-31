#!/usr/bin/env python3
"""Generate one-page lean methodology integration PPT for PV manufacturing."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsmap
from lxml import etree

# Visual direction: industrial navy + solar teal (avoid purple / cream / broadsheet)
NAVY = RGBColor(0x0B, 0x2A, 0x3F)
NAVY_MID = RGBColor(0x14, 0x3D, 0x58)
TEAL = RGBColor(0x0D, 0x8A, 0x7A)
TEAL_LIGHT = RGBColor(0xE6, 0xF5, 0xF2)
SLATE = RGBColor(0x2C, 0x3E, 0x50)
GRAY = RGBColor(0x5A, 0x6A, 0x7A)
LIGHT_BG = RGBColor(0xF4, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_GOLD = RGBColor(0xC4, 0x8A, 0x2A)
SOFT_LINE = RGBColor(0xD0, 0xDB, 0xE4)
CARD_BG = RGBColor(0xFF, 0xFF, 0xFF)
LEVEL_COLORS = [
    RGBColor(0x0B, 0x2A, 0x3F),  # HQ
    RGBColor(0x0D, 0x6B, 0x8A),  # Base
    RGBColor(0x0D, 0x8A, 0x7A),  # Workshop
    RGBColor(0x2E, 0x8B, 0x57),  # Team
]


def set_run_font(run, name="Microsoft YaHei", size=10, bold=False, color=SLATE):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # East Asian font fallback
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    ea.set("typeface", name)


def add_textbox(slide, left, top, width, height, text, size=10, bold=False,
                color=SLATE, align=PP_ALIGN.LEFT, font="Microsoft YaHei",
                anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {
            MSO_ANCHOR.TOP: "t",
            MSO_ANCHOR.MIDDLE: "ctr",
            MSO_ANCHOR.BOTTOM: "b",
        }.get(anchor, "t"))
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    set_run_font(run, name=font, size=size, bold=bold, color=color)
    return box


def add_rich_textbox(slide, left, top, width, height, paragraphs, anchor=MSO_ANCHOR.TOP):
    """paragraphs: list of dicts with text, size, bold, color, align, space_after."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set("anchor", {
            MSO_ANCHOR.TOP: "t",
            MSO_ANCHOR.MIDDLE: "ctr",
            MSO_ANCHOR.BOTTOM: "b",
        }.get(anchor, "t"))
    except Exception:
        pass
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", PP_ALIGN.LEFT)
        p.space_before = Pt(para.get("space_before", 0))
        p.space_after = Pt(para.get("space_after", 2))
        run = p.add_run()
        run.text = para["text"]
        set_run_font(
            run,
            name=para.get("font", "Microsoft YaHei"),
            size=para.get("size", 10),
            bold=para.get("bold", False),
            color=para.get("color", SLATE),
        )
    return box


def add_rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    return shape


def add_rounded(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    # Soften corner
    try:
        adj = shape.adjustments
        adj[0] = 0.08
    except Exception:
        pass
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    return shape


def add_chevron(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # ===== Background =====
    add_rect(slide, 0, 0, prs.slide_width, prs.slide_height, LIGHT_BG)
    # Top brand bar
    add_rect(slide, 0, 0, prs.slide_width, Inches(0.92), NAVY)
    # Accent strip under header
    add_rect(slide, 0, Inches(0.92), prs.slide_width, Inches(0.06), TEAL)
    # Bottom bar
    add_rect(slide, 0, Inches(7.15), prs.slide_width, Inches(0.35), NAVY)

    # ===== Title =====
    add_textbox(
        slide, Inches(0.35), Inches(0.12), Inches(11.5), Inches(0.42),
        "精益工具方法论系统化整合｜四级联动落地体系（含广州组件基地）",
        size=20, bold=True, color=WHITE, align=PP_ALIGN.LEFT,
    )
    add_textbox(
        slide, Inches(0.35), Inches(0.52), Inches(10.5), Inches(0.32),
        "适配硅片 / 组件全产线  ·  统一术语 · 统一推行流程 · 统一考核标准  ·  真善美价值主张贯穿",
        size=11, bold=False, color=RGBColor(0xA8, 0xC5, 0xD4),
    )

    # ===== 1. Core goals strip =====
    y0 = Inches(1.12)
    add_rounded(slide, Inches(0.28), y0, Inches(12.77), Inches(0.58), WHITE, SOFT_LINE)
    add_rect(slide, Inches(0.28), y0, Inches(0.12), Inches(0.58), TEAL)
    add_textbox(
        slide, Inches(0.55), y0 + Inches(0.06), Inches(1.6), Inches(0.22),
        "核心目标", size=12, bold=True, color=TEAL,
    )
    goals = ["提质", "降损耗", "提设备效率", "均衡多基地产能", "低碳降耗保交付"]
    gx = 0.55
    for i, g in enumerate(goals):
        left = Inches(gx + i * 2.35)
        add_rounded(slide, left, y0 + Inches(0.28), Inches(2.15), Inches(0.24), TEAL_LIGHT)
        add_textbox(
            slide, left, y0 + Inches(0.28), Inches(2.15), Inches(0.24),
            g, size=10, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )

    # ===== Section headers helper =====
    def section_title(x, y, w, title):
        add_textbox(slide, x, y, w, Inches(0.28), title, size=13, bold=True, color=NAVY)

    # ===== 2. Five modules =====
    section_title(Inches(0.28), Inches(1.82), Inches(6), "精益工具五大集成模块")

    modules = [
        ("① 价值流诊断", "VSM价值流图\n八大浪费识别\n产销协同均衡", NAVY),
        ("② 现场精益", "6S / 目视化\nSMED快速换型\n线平衡 · Poka-Yoke", RGBColor(0x0D, 0x6B, 0x8A)),
        ("③ 品质管控", "5Why / 鱼骨图\n柏拉图 · SPC · MSA\n8D · OCAP闭环", TEAL),
        ("④ 设备能耗", "TPM全员保全\nOEE三维拆解\nMTTR/MTBF · 能源精益", RGBColor(0x1A, 0x7A, 0x5C)),
        ("⑤ 长效固化", "SOP / 标准工时\n全员Kaizen\n精益数字化看板", RGBColor(0x2E, 0x8B, 0x57)),
    ]

    card_w = Inches(2.42)
    card_h = Inches(1.85)
    gap = Inches(0.12)
    start_x = Inches(0.28)
    my = Inches(2.12)

    for i, (title, body, color) in enumerate(modules):
        x = start_x + i * (card_w + gap)
        add_rounded(slide, x, my, card_w, card_h, WHITE, SOFT_LINE)
        add_rect(slide, x, my, card_w, Inches(0.38), color)
        add_textbox(
            slide, x + Inches(0.08), my + Inches(0.05), card_w - Inches(0.16), Inches(0.28),
            title, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
        )
        lines = body.split("\n")
        paras = []
        for j, line in enumerate(lines):
            paras.append({
                "text": line,
                "size": 11,
                "bold": False,
                "color": SLATE,
                "align": PP_ALIGN.CENTER,
                "space_after": 6 if j < len(lines) - 1 else 0,
                "space_before": 4 if j == 0 else 0,
            })
        add_rich_textbox(
            slide, x + Inches(0.1), my + Inches(0.48), card_w - Inches(0.2), Inches(1.25),
            paras, anchor=MSO_ANCHOR.MIDDLE,
        )

    # ===== 3. Four-level linkage =====
    section_title(Inches(0.28), Inches(4.12), Inches(8), "四级联动推行架构（覆盖广州组件基地）")

    levels = [
        ("一级 · 总部精益部", "标准制定 · 跨基地VSM统筹\n数字化平台 · 考核指标", "珠海 · 西宁 · 宜宾 · 广州"),
        ("二级 · 基地管理部", "基地VSM · TPM/5S专项\n成果汇总 · 跨车间协调", "含广州组件基地"),
        ("三级 · 车间主任", "线平衡 · SMED · 8D\n车间OEE与损耗管控", "硅片 / 组件产线"),
        ("四级 · 班组 / 线长", "5S目视化 · 工位5Why\n防错改善 · Kaizen提案", "一线高频落地"),
    ]

    ly = Inches(4.42)
    lw = Inches(3.10)
    lh = Inches(1.35)
    lg = Inches(0.12)

    for i, (title, body, tag) in enumerate(levels):
        x = Inches(0.28) + i * (lw + lg)
        add_rounded(slide, x, ly, lw, lh, WHITE, SOFT_LINE)
        add_rect(slide, x, ly, lw, Inches(0.36), LEVEL_COLORS[i])
        add_textbox(
            slide, x + Inches(0.08), ly + Inches(0.05), lw - Inches(0.16), Inches(0.28),
            title, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
        )
        paras = []
        for j, line in enumerate(body.split("\n")):
            paras.append({
                "text": line,
                "size": 10,
                "bold": False,
                "color": SLATE,
                "align": PP_ALIGN.CENTER,
                "space_after": 3,
                "space_before": 2 if j == 0 else 0,
            })
        paras.append({
            "text": tag,
            "size": 9,
            "bold": True,
            "color": TEAL,
            "align": PP_ALIGN.CENTER,
            "space_before": 4,
            "space_after": 0,
        })
        add_rich_textbox(
            slide, x + Inches(0.08), ly + Inches(0.42), lw - Inches(0.16), Inches(0.88),
            paras,
        )
        # Arrow between levels
        if i < 3:
            ax = x + lw + Inches(0.01)
            add_textbox(
                slide, ax, ly + Inches(0.5), Inches(0.12), Inches(0.3),
                "→", size=14, bold=True, color=TEAL, align=PP_ALIGN.CENTER,
            )

    # ===== 4. Closed-loop process =====
    cy = Inches(5.95)
    add_textbox(
        slide, Inches(0.28), cy, Inches(2.2), Inches(0.28),
        "统一闭环机制", size=13, bold=True, color=NAVY,
    )

    steps = ["现状诊断", "目标拆解", "工具落地", "数据验证", "标准化固化", "跨基地复制", "持续迭代"]
    sx = Inches(2.35)
    sw = Inches(1.48)
    sg = Inches(0.02)
    for i, step in enumerate(steps):
        x = sx + i * (sw + sg)
        color = TEAL if i % 2 == 0 else NAVY_MID
        chev = add_chevron(slide, x, cy + Inches(0.02), sw, Inches(0.42), color)
        # Text on chevron — use text frame of shape
        tf = chev.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = step
        set_run_font(run, size=10, bold=True, color=WHITE)

    # Note under loop
    add_textbox(
        slide, Inches(2.35), cy + Inches(0.48), Inches(10.5), Inches(0.22),
        "广州组件成功经验同步至珠海 / 西宁 / 宜宾等基地  ·  指标四级逐级分解  ·  杜绝工具零散、重复落地",
        size=9, color=GRAY,
    )

    # Footer
    add_textbox(
        slide, Inches(0.35), Inches(7.18), Inches(8), Inches(0.28),
        "精益五星模型：价值流诊断 → 现场基础管理 → 品质闭环 → 设备能耗优化 → 持续改善长效机制",
        size=9, color=RGBColor(0xA8, 0xC5, 0xD4),
    )
    add_textbox(
        slide, Inches(9.5), Inches(7.18), Inches(3.5), Inches(0.28),
        "硅片 · 组件 · 广州基地四级联动",
        size=9, bold=True, color=WHITE, align=PP_ALIGN.RIGHT,
    )

    from pathlib import Path
    out_dir = Path(__file__).resolve().parent / "output"
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "精益工具方法论系统化整合_四级联动一页PPT.pptx"
    prs.save(str(out))
    print(f"Saved: {out}")
    return str(out)


if __name__ == "__main__":
    build()
